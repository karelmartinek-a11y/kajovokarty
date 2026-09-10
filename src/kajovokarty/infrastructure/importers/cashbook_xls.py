from __future__ import annotations

import csv
import json
import unicodedata
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from xlrd import open_workbook

from ...domain.enums import MatchStatus, RunState
from ..persistence.database import Database
from .common import (
    ImportCancelled,
    ImportValidationError,
    immutable_snapshot,
    money_minor,
    normalize_text,
    parse_datetime,
    sha256_text,
    canonical_json,
    utc_now,
)

CASHBOOK_HEADERS = (
    "Vystaveno", "Pohyb", "Číslo", "Označení", "Klient", "Příjem", "Výdaj",
    "Stav pokladny", "Měna", "Forma úhrady", "Pokladna", "Variabilní symbol", "Vystavil",
)
CARD_FORMS = {"card", "karta", "platebni karta", "platební karta", "payment card"}
SYSTEM_ROWS = {"počáteční stav", "pocatecni stav", "uzávěrka", "uzaverka", "konečný stav", "konecny stav"}
def _prague_timezone():
    try:
        return ZoneInfo("Europe/Prague")
    except Exception:
        # Minimal installations may omit tzdata; keep imports usable and use
        # the standard winter offset until the host supplies the timezone DB.
        return timezone(timedelta(hours=1))


def _header_key(value: object) -> str:
    text = unicodedata.normalize("NFKC", normalize_text(value)).casefold()
    return "".join(ch for ch in unicodedata.normalize("NFD", text) if unicodedata.category(ch) != "Mn")


@dataclass(frozen=True, slots=True)
class CashbookRow:
    row_no: int
    raw: dict[str, object]
    identity: str | None
    content_hash: str | None
    occurred_at: str | None
    amount_minor: int = 0
    currency: str = ""
    direction: str = ""
    relevant: bool = False
    reason: str | None = None


@dataclass(frozen=True, slots=True)
class CashbookSheet:
    name: str
    rows: list[list[object]]


@dataclass(slots=True)
class CashbookImportReport:
    run_id: str
    correlation_id: str
    file_hash: str
    rows_seen: int = 0
    new_rows: int = 0
    duplicate_rows: int = 0
    conflicts: int = 0
    ignored_rows: int = 0
    totals_minor: dict[str, int] = field(default_factory=dict)
    state: RunState = RunState.PENDING


class CashbookFileImportService:
    """Import Better Hotel card movements as immutable reconciliation sources."""

    def __init__(self, database: Database) -> None:
        self.database = database

    def import_file(self, path: Path, *, cancel_callback=None, correlation_id: str | None = None) -> CashbookImportReport:
        snapshot = immutable_snapshot(path)
        run_id = uuid4().hex
        correlation = correlation_id or uuid4().hex
        report = CashbookImportReport(run_id, correlation, snapshot.file_hash, state=RunState.RUNNING)
        started = utc_now()
        try:
            sheets = self._load_sheets(snapshot.snapshot_path)
            sheet = self._select_sheet(sheets)
            rows = self._parse_sheet(sheet.rows)
            report.rows_seen = len(rows)
            with self.database.transaction() as conn:
                conn.execute(
                    "INSERT INTO cashbook_import_run(id,correlation_id,file_name,file_hash,state,counts_json,totals_json,started_at_utc) VALUES(?,?,?,?,?,?,?,?)",
                    (run_id, correlation, path.name, snapshot.file_hash, RunState.RUNNING.value, "{}", "{}", started),
                )
                for row in rows:
                    if cancel_callback is not None and cancel_callback():
                        raise ImportCancelled("Operace byla bezpečně zrušena.")
                    if not row.relevant:
                        report.ignored_rows += 1
                        continue
                    existing = conn.execute(
                        "SELECT content_hash FROM cashbook_card_transaction WHERE cashbook_identity=?",
                        (row.identity,),
                    ).fetchone()
                    if existing is not None:
                        if existing["content_hash"] == row.content_hash:
                            report.duplicate_rows += 1
                            continue
                        report.conflicts += 1
                        raise ImportValidationError(
                            f"Konflikt pokladního dokladu {row.identity} na řádku {row.row_no}; import byl vrácen."
                        )
                    now = utc_now()
                    conn.execute(
                        "INSERT INTO cashbook_card_transaction(id,cashbook_identity,occurred_at,direction,amount_minor,currency_code,receipt_number,label,client,payment_form,cashbox,variable_symbol,issued_by,status,raw_json,content_hash,first_seen_utc,last_seen_utc) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (
                            uuid4().hex, row.identity, row.occurred_at, row.direction, row.amount_minor, row.currency,
                            row.identity, normalize_text(row.raw["Označení"]), normalize_text(row.raw["Klient"]) or None,
                            normalize_text(row.raw["Forma úhrady"]), normalize_text(row.raw["Pokladna"]) or None,
                            normalize_text(row.raw["Variabilní symbol"]) or None, normalize_text(row.raw["Vystavil"]) or None,
                            MatchStatus.UNMATCHED.value, canonical_json(row.raw), row.content_hash, now, now,
                        ),
                    )
                    report.new_rows += 1
                    report.totals_minor[row.currency] = report.totals_minor.get(row.currency, 0) + row.amount_minor
                report.state = RunState.SUCCEEDED
                conn.execute(
                    "UPDATE cashbook_import_run SET state=?,counts_json=?,totals_json=?,finished_at_utc=? WHERE id=?",
                    (report.state.value, canonical_json(self._counts(report)), canonical_json(report.totals_minor), utc_now(), run_id),
                )
            return report
        except ImportCancelled:
            report.state = RunState.CANCELLED
            self._record_failure(report, path.name, started, "cancelled")
            raise
        except Exception as exc:
            report.state = RunState.FAILED
            self._record_failure(report, path.name, started, str(exc))
            raise
        finally:
            snapshot.cleanup()

    @staticmethod
    def _counts(report: CashbookImportReport) -> dict[str, int]:
        return {"rows_seen": report.rows_seen, "new": report.new_rows, "duplicate": report.duplicate_rows, "conflict": report.conflicts, "ignored": report.ignored_rows}

    def _record_failure(self, report: CashbookImportReport, file_name: str, started: str, error: str) -> None:
        with self.database.transaction() as conn:
            conn.execute(
                "INSERT INTO cashbook_import_run(id,correlation_id,file_name,file_hash,state,counts_json,totals_json,error_json,started_at_utc,finished_at_utc) VALUES(?,?,?,?,?,?,?,?,?,?)",
                (report.run_id, report.correlation_id, file_name, report.file_hash, report.state.value, canonical_json(self._counts(report)), canonical_json(report.totals_minor), canonical_json({"message": error}), started, utc_now()),
            )

    @staticmethod
    def _load_sheets(path: Path) -> list[CashbookSheet]:
        if path.suffix.casefold() == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                return [CashbookSheet(path.stem, [list(row) for row in csv.reader(handle)])]
        if path.suffix.casefold() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                return [CashbookSheet(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]) for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
            finally:
                workbook.close()
        if path.suffix.casefold() == ".xls":
            workbook = open_workbook(path, on_demand=True, formatting_info=False)
            try:
                return [CashbookSheet(sheet.name, [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]) for sheet in (workbook.sheet_by_index(i) for i in range(workbook.nsheets)) if sheet.visibility == 0]
            finally:
                workbook.release_resources()
        raise ImportValidationError("Pokladní deník podporuje pouze XLS, XLSX a CSV.")

    @staticmethod
    def _select_sheet(sheets: list[CashbookSheet]) -> CashbookSheet:
        matches = [sheet for sheet in sheets if any(CashbookFileImportService._is_header(row) for row in sheet.rows[:50])]
        if len(matches) != 1:
            raise ImportValidationError("Soubor musí obsahovat právě jeden viditelný list s hlavičkou pokladního deníku.")
        return matches[0]

    @staticmethod
    def _is_header(row: list[object]) -> bool:
        keys = {_header_key(value) for value in row}
        return all(_header_key(header) in keys for header in CASHBOOK_HEADERS)

    def _parse_sheet(self, sheet: list[list[object]]) -> list[CashbookRow]:
        header_row = next((row for row in sheet[:50] if self._is_header(row)), None)
        if header_row is None:
            raise ImportValidationError("Hlavička pokladního deníku nebyla nalezena.")
        header_index = sheet.index(header_row)
        header = {_header_key(value): index for index, value in enumerate(header_row)}
        result: list[CashbookRow] = []
        for row_no, values in enumerate(sheet[header_index + 1 :], start=header_index + 2):
            raw = {header_key: values[index] if index < len(values) else None for header_key, index in header.items()}
            named = {header: raw[_header_key(header)] for header in CASHBOOK_HEADERS}
            if not any(normalize_text(value) for value in values):
                result.append(CashbookRow(row_no, named, None, None, None, reason="BLANK"))
                continue
            label = normalize_text(named["Označení"]).casefold()
            if label in SYSTEM_ROWS or not normalize_text(named["Číslo"]):
                result.append(CashbookRow(row_no, named, None, None, None, reason="SYSTEM_OR_NO_ID"))
                continue
            form = normalize_text(named["Forma úhrady"]).casefold()
            if form not in CARD_FORMS:
                result.append(CashbookRow(row_no, named, None, None, None, reason="NON_CARD"))
                continue
            income, currency = money_minor(named["Příjem"] or 0, named["Měna"])
            expense, _ = money_minor(named["Výdaj"] or 0, currency)
            if income and expense:
                raise ImportValidationError(f"Řádek {row_no} má současně příjem i výdaj.")
            if not income and not expense:
                result.append(CashbookRow(row_no, named, None, None, None, reason="ZERO"))
                continue
            amount = income if income else -expense
            timestamp = parse_datetime(named["Vystaveno"])
            assert timestamp is not None
            if timestamp.tzinfo is None:
                timestamp = timestamp.replace(tzinfo=_prague_timezone())
            occurred = timestamp.astimezone(UTC).isoformat(timespec="seconds").replace("+00:00", "Z")
            identity = normalize_text(named["Číslo"])
            if identity.endswith(".0") and identity[:-2].isdigit():
                identity = identity[:-2]
            canonical = {key: normalize_text(value) for key, value in named.items()}
            canonical.update({"direction": "INCOME" if amount > 0 else "EXPENSE", "amount_minor": amount, "currency": currency, "occurred_at": occurred, "identity": identity})
            result.append(CashbookRow(row_no, named, identity, sha256_text(canonical_json(canonical)), occurred, amount, currency, "INCOME" if amount > 0 else "EXPENSE", True))
        return result
