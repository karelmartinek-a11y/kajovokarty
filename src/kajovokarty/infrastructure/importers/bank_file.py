from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from xlrd import open_workbook

from ...domain.entities import CardTransaction
from ...domain.enums import MatchStatus, RunState
from ..persistence.database import Database
from .common import (
    CancelCallback,
    ImportCancelled,
    ImportValidationError,
    ProgressCallback,
    canonical_json,
    cancellation_point,
    immutable_snapshot,
    map_headers,
    money_minor,
    normalize_header,
    normalize_text,
    parse_czech_date,
    parse_datetime,
    progress,
    row_dict,
    sha256_text,
    utc_now,
)

BANK_HEADERS = (
    "Typ transakce",
    "ID Terminálu",
    "ID POS",
    "Datum a čas vzniku",
    "Čas připsání na server",
    "Datum zaúčtování",
    "Částka",
    "Cashback",
    "Spropitné",
    "Měna",
    "ARN kód",
    "DCC",
    "Číslo karty/Číslo účtu",
    "Autoriz. kód",
    "Var. symbol",
    "Var. symbol 2",
    "SEQ ID",
    "Vydavatel karty",
    "Způsob načtení karty",
    "Obchodní místo",
    "Adresa obchodního místa",
)

SALE_TYPES = {"prodej"}
CLOSURE_TYPES = {"uzávěrka", "uzaverka"}
REFUND_TYPES = {"refundace", "storno", "vrácení", "vraceni", "refund"}
SUMMARY_PREFIXES = ("počet transakcí:", "pocet transakci:", "suma částek:", "suma castek:")


@dataclass(frozen=True, slots=True)
class SheetData:
    name: str
    rows: list[list[object]]


@dataclass(frozen=True, slots=True)
class ParsedBankRow:
    row_no: int
    kind: str
    raw: dict[str, object]
    transaction: CardTransaction | None
    reason: str | None = None
    source_summary: str | None = None


@dataclass(slots=True)
class BankImportReport:
    run_id: str
    correlation_id: str
    file_hash: str
    selected_sheet: str = ""
    rows_seen: int = 0
    new_rows: int = 0
    duplicate_rows: int = 0
    conflicts: int = 0
    sales: int = 0
    closures: int = 0
    blank_rows: int = 0
    summary_rows: int = 0
    quarantined: int = 0
    totals_minor: dict[str, int] = field(default_factory=dict)
    state: RunState = RunState.PENDING


class MultipleMatchingSheets(ImportValidationError):
    def __init__(self, names: list[str]) -> None:
        super().__init__("Více listů odpovídá bankovnímu schématu; zvolte list v UI: " + ", ".join(names))
        self.names = names


class BankFileImportService:
    def __init__(self, database: Database) -> None:
        self.database = database

    def inspect(self, path: Path, *, sheet_name: str | None = None) -> tuple[str, list[ParsedBankRow]]:
        snapshot = immutable_snapshot(path)
        try:
            sheet = self._select_sheet(self._load_sheets(snapshot.snapshot_path), sheet_name)
            return sheet.name, self._parse_sheet(sheet, self._load_mappings())
        finally:
            snapshot.cleanup()

    def import_file(
        self,
        path: Path,
        *,
        sheet_name: str | None = None,
        progress_callback: ProgressCallback | None = None,
        cancel_callback: CancelCallback | None = None,
        correlation_id: str | None = None,
    ) -> BankImportReport:
        snapshot = immutable_snapshot(path)
        run_id = uuid4().hex
        correlation = correlation_id or uuid4().hex
        report = BankImportReport(run_id, correlation, snapshot.file_hash, state=RunState.RUNNING)
        started = utc_now()
        try:
            previous = self.database.query(
                "SELECT counts_json,totals_json,source_summary_json FROM bank_import_run WHERE file_hash=? AND state IN ('SUCCEEDED','DUPLICATE') ORDER BY finished_at_utc DESC LIMIT 1",
                (snapshot.file_hash,),
            )
            if previous:
                counts = json.loads(previous[0]["counts_json"] or "{}")
                report.rows_seen = int(counts.get("rows_seen", 0))
                report.new_rows = 0
                report.duplicate_rows = int(counts.get("sales", counts.get("new", 0)))
                report.sales = int(counts.get("sales", 0))
                report.closures = int(counts.get("closures", 0))
                report.blank_rows = int(counts.get("blank", 0))
                report.summary_rows = int(counts.get("summary", 0))
                report.totals_minor = {str(k): int(v) for k, v in json.loads(previous[0]["totals_json"] or "{}").items()}
                summary = json.loads(previous[0]["source_summary_json"] or "{}")
                report.selected_sheet = str(summary.get("sheet", ""))
                report.state = RunState.DUPLICATE
                self._record_duplicate(path, report, started)
                return report
            sheet = self._select_sheet(self._load_sheets(snapshot.snapshot_path), sheet_name)
            report.selected_sheet = sheet.name
            parsed_rows = self._parse_sheet(sheet, self._load_mappings())
            report.rows_seen = len(sheet.rows)
            total = len(parsed_rows)
            # The operation progress callback writes to SQLite. Keep it outside
            # the import transaction to avoid a second writer deadlocking the
            # transaction that is processing the rows.
            progress(progress_callback, 0, total, f"Banka: připravuji {total} řádků")
            with self.database.transaction() as conn:
                conn.execute(
                    "INSERT INTO bank_import_run(id,correlation_id,file_name,file_hash,state,counts_json,totals_json,source_summary_json,started_at_utc,heartbeat_at_utc) VALUES(?,?,?,?,?,?,?,?,?,?)",
                    (run_id, correlation, path.name, snapshot.file_hash, RunState.RUNNING.value, "{}", "{}", "{}", started, started),
                )
                for index, parsed in enumerate(parsed_rows, start=1):
                    cancellation_point(cancel_callback)
                    self._apply_row(conn, run_id, parsed, report)
                    if index % 20 == 0:
                        conn.execute("UPDATE bank_import_run SET heartbeat_at_utc=? WHERE id=?", (utc_now(), run_id))
                report.state = RunState.SUCCEEDED
                conn.execute(
                    "UPDATE bank_import_run SET state=?,counts_json=?,totals_json=?,source_summary_json=?,heartbeat_at_utc=?,finished_at_utc=? WHERE id=?",
                    (
                        report.state.value,
                        canonical_json(self._counts(report)),
                        canonical_json(report.totals_minor),
                        canonical_json({"sheet": report.selected_sheet}),
                        utc_now(),
                        utc_now(),
                        run_id,
                    ),
                )
            progress(progress_callback, total, total, f"Banka: hotovo, {total} řádků")
            return report
        except ImportCancelled:
            report.state = RunState.CANCELLED
            self._record_failed(path, report, started, "cancelled")
            raise
        except Exception as exc:
            report.state = RunState.FAILED
            self._record_failed(path, report, started, str(exc))
            raise
        finally:
            snapshot.cleanup()

    @staticmethod
    def _counts(report: BankImportReport) -> dict[str, int]:
        return {
            "rows_seen": report.rows_seen,
            "new": report.new_rows,
            "duplicate": report.duplicate_rows,
            "conflict": report.conflicts,
            "sales": report.sales,
            "closures": report.closures,
            "blank": report.blank_rows,
            "summary": report.summary_rows,
            "quarantine": report.quarantined,
        }

    def _record_failed(self, path: Path, report: BankImportReport, started: str, error: str) -> None:
        with self.database.transaction() as conn:
            conn.execute(
                "INSERT INTO bank_import_run(id,correlation_id,file_name,file_hash,state,counts_json,totals_json,source_summary_json,error_json,started_at_utc,heartbeat_at_utc,finished_at_utc) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET state=excluded.state,error_json=excluded.error_json,heartbeat_at_utc=excluded.heartbeat_at_utc,finished_at_utc=excluded.finished_at_utc",
                (
                    report.run_id,
                    report.correlation_id,
                    path.name,
                    report.file_hash,
                    report.state.value,
                    canonical_json(self._counts(report)),
                    canonical_json(report.totals_minor),
                    canonical_json({"sheet": report.selected_sheet}),
                    canonical_json({"message": error}),
                    started,
                    utc_now(),
                    utc_now(),
                ),
            )

    def _record_duplicate(self, path: Path, report: BankImportReport, started: str) -> None:
        with self.database.transaction() as conn:
            conn.execute(
                "INSERT INTO bank_import_run(id,correlation_id,file_name,file_hash,state,counts_json,totals_json,source_summary_json,error_json,started_at_utc,heartbeat_at_utc,finished_at_utc) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                (report.run_id, report.correlation_id, path.name, report.file_hash, report.state.value, canonical_json(self._counts(report)), canonical_json(report.totals_minor), canonical_json({"sheet": report.selected_sheet, "message": "Přesný duplikát souboru byl přeskočen."}), None, started, utc_now(), utc_now()),
            )

    def _load_mappings(self) -> dict[str, str]:
        return {
            normalize_text(row["raw_value"]).casefold(): normalize_text(row["mapped_meaning"]).upper()
            for row in self.database.query("SELECT raw_value,mapped_meaning FROM import_value_mapping WHERE source_kind='BANK' AND field_name='transaction_type'")
        }

    def reprocess_quarantine(self, quarantine_id: str) -> BankImportReport:
        rows = self.database.query("SELECT * FROM quarantined_source_row WHERE id=? AND run_type='BANK'", (quarantine_id,))
        if not rows:
            raise ImportValidationError("Karanténní bankovní řádek neexistuje.")
        quarantine = rows[0]
        raw = json.loads(quarantine["raw_json"])
        parsed = self._parse_raw_row(int(quarantine["row_no"]), raw, self._load_mappings())
        if parsed.kind == "quarantine":
            self.database.execute("UPDATE quarantined_source_row SET retry_count=retry_count+1,updated_at_utc=? WHERE id=?", (utc_now(), quarantine_id))
            raise ImportValidationError("Řádek stále nemá bezpečně namapovaný typ transakce.")
        report = BankImportReport(str(quarantine["run_id"]), uuid4().hex, "reprocess", rows_seen=1, state=RunState.RUNNING)
        with self.database.transaction() as conn:
            self._apply_row(conn, str(quarantine["run_id"]), parsed, report)
            if report.quarantined or report.conflicts:
                conn.execute("UPDATE quarantined_source_row SET retry_count=retry_count+1,updated_at_utc=? WHERE id=?", (utc_now(), quarantine_id))
                raise ImportValidationError("Řádek nelze bezpečně znovu zpracovat; vznikl konflikt.")
            conn.execute("UPDATE quarantined_source_row SET state='RESOLVED',resolution_method='REPROCESSED',resolution_json=?,resolved_at_utc=?,updated_at_utc=?,retry_count=retry_count+1 WHERE id=?", (canonical_json({"new": report.new_rows, "duplicate": report.duplicate_rows}), utc_now(), utc_now(), quarantine_id))
        report.state = RunState.SUCCEEDED
        return report

    def _load_sheets(self, path: Path) -> list[SheetData]:
        suffix = path.suffix.casefold()
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = [list(row) for row in csv.reader(handle, dialect="excel", strict=True)]
            return [SheetData(path.stem, rows)]
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                return [
                    SheetData(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
                    for sheet in workbook.worksheets
                    if sheet.sheet_state == "visible"
                ]
            finally:
                workbook.close()
        if suffix == ".xls":
            workbook = open_workbook(path, on_demand=True, formatting_info=False)
            result: list[SheetData] = []
            try:
                for index in range(workbook.nsheets):
                    sheet = workbook.sheet_by_index(index)
                    visibility = getattr(workbook, "sheet_visibility", [0] * workbook.nsheets)[index]
                    if visibility != 0:
                        continue
                    rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
                    result.append(SheetData(sheet.name, rows))
            finally:
                workbook.release_resources()
            return result
        raise ImportValidationError("Podporované přípony bankovního souboru jsou CSV, XLS a XLSX.")

    def _select_sheet(self, sheets: list[SheetData], requested: str | None) -> SheetData:
        matches: list[SheetData] = []
        for sheet in sheets:
            if any(self._is_header_row(row) for row in sheet.rows[:50]):
                matches.append(sheet)
        if requested is not None:
            for sheet in matches:
                if sheet.name == requested:
                    return sheet
            raise ImportValidationError(f"Zvolený list {requested!r} neodpovídá očekávanému schématu.")
        if not matches:
            raise ImportValidationError("V souboru nebyl nalezen viditelný list s bankovní hlavičkou.")
        if len(matches) > 1:
            raise MultipleMatchingSheets([sheet.name for sheet in matches])
        return matches[0]

    @staticmethod
    def _is_header_row(row: list[object]) -> bool:
        normalized = {normalize_header(cell) for cell in row}
        return all(normalize_header(required) in normalized for required in BANK_HEADERS)

    def _parse_sheet(self, sheet: SheetData, mappings: dict[str, str] | None = None) -> list[ParsedBankRow]:
        header_index = next((i for i, row in enumerate(sheet.rows) if self._is_header_row(row)), None)
        if header_index is None:
            raise ImportValidationError("Bankovní hlavička nebyla nalezena.")
        mapping = map_headers(sheet.rows[header_index], BANK_HEADERS)
        mappings = mappings or {}
        result: list[ParsedBankRow] = []
        for row_no, row in enumerate(sheet.rows[header_index + 1 :], start=header_index + 2):
            raw = row_dict(row, mapping)
            result.append(self._parse_raw_row(row_no, raw, mappings, is_blank=not any(normalize_text(value) for value in row)))
        return result

    def _parse_raw_row(self, row_no: int, raw: dict[str, object], mappings: dict[str, str] | None = None, *, is_blank: bool = False) -> ParsedBankRow:
        mappings = mappings or {}
        if is_blank or not any(normalize_text(value) for value in raw.values()):
            return ParsedBankRow(row_no, "blank", raw, None)
        row_type = normalize_text(raw["Typ transakce"])
        folded = row_type.casefold()
        mapped = mappings.get(folded)
        if mapped == "SALE":
            folded = "prodej"
        elif mapped == "REFUND":
            folded = "refundace"
        elif mapped == "CLOSURE":
            folded = "uzávěrka"
        elif mapped == "IGNORE":
            return ParsedBankRow(row_no, "summary", raw, None, source_summary=f"Vědomě ignorováno: {row_type}")
        if any(folded.startswith(prefix) for prefix in SUMMARY_PREFIXES):
            return ParsedBankRow(row_no, "summary", raw, None, source_summary=row_type)
        if folded in CLOSURE_TYPES:
            return ParsedBankRow(row_no, "closure", raw, None)
        if folded not in SALE_TYPES and folded not in REFUND_TYPES:
            return ParsedBankRow(row_no, "quarantine", raw, None, "UNKNOWN_TRANSACTION_TYPE")
        try:
            return self._parse_transaction(row_no, raw, folded in REFUND_TYPES)
        except ImportValidationError as exc:
            return ParsedBankRow(row_no, "quarantine", raw, None, f"ROW_PARSE_ERROR:{exc}")

    def _parse_transaction(self, row_no: int, raw: dict[str, object], refund: bool) -> ParsedBankRow:
        terminal_id = normalize_text(raw["ID Terminálu"])
        seq_id = normalize_text(raw["SEQ ID"])
        if not terminal_id or not seq_id:
            raise ImportValidationError("ID Terminálu nebo SEQ ID chybí.")
        amount, currency = money_minor(raw["Částka"], raw["Měna"])
        if refund and amount > 0:
            amount = -amount
        if amount == 0:
            return ParsedBankRow(row_no, "quarantine", raw, None, "ZERO_AMOUNT")
        occurred = parse_datetime(raw["Datum a čas vzniku"])
        server_at = parse_datetime(raw["Čas připsání na server"], allow_empty=True)
        posted = parse_czech_date(raw["Datum zaúčtování"], allow_empty=True)
        cashback = self._optional_money(raw["Cashback"], currency)
        tip = self._optional_money(raw["Spropitné"], currency)
        dcc = self._optional_money(raw["DCC"], currency)
        transaction_type = normalize_text(raw["Typ transakce"])
        content_hash = sha256_text(canonical_json(raw))
        transaction = CardTransaction(
            terminal_id=terminal_id,
            seq_id=seq_id,
            amount_minor=amount,
            currency=currency,
            occurred_at=occurred,
            posted_date=posted,
            transaction_type=transaction_type,
            server_at=server_at,
            arn=normalize_text(raw["ARN kód"]) or None,
            authorization_code=normalize_text(raw["Autoriz. kód"]) or None,
            masked_card=normalize_text(raw["Číslo karty/Číslo účtu"]) or None,
            variable_symbol=normalize_text(raw["Var. symbol"]) or None,
            variable_symbol_2=normalize_text(raw["Var. symbol 2"]) or None,
            cashback_minor=cashback,
            tip_minor=tip,
            dcc_minor=dcc,
            status=MatchStatus.UNMATCHED,
            content_hash=content_hash,
        )
        return ParsedBankRow(row_no, "sale", raw, transaction)

    @staticmethod
    def _optional_money(value: object, currency: str) -> int:
        if value is None or normalize_text(value) == "":
            return 0
        amount, _ = money_minor(value, currency)
        return amount

    def _apply_row(self, conn: Any, run_id: str, parsed: ParsedBankRow, report: BankImportReport) -> None:
        if parsed.kind == "blank":
            report.blank_rows += 1
            return
        if parsed.kind == "summary":
            report.summary_rows += 1
            return
        if parsed.kind == "closure":
            report.closures += 1
            return
        if parsed.kind == "quarantine" or parsed.transaction is None:
            report.quarantined += 1
            self._quarantine(conn, run_id, parsed)
            return
        transaction = parsed.transaction
        existing = conn.execute(
            "SELECT id,content_hash FROM card_transaction WHERE terminal_id=? AND seq_id=?",
            (transaction.terminal_id, transaction.seq_id),
        ).fetchone()
        if existing is not None:
            if existing["content_hash"] == transaction.content_hash:
                report.duplicate_rows += 1
                conn.execute(
                    "UPDATE card_transaction SET last_seen_utc=? WHERE id=?", (utc_now(), existing["id"])
                )
                return
            report.conflicts += 1
            self._quarantine(conn, run_id, ParsedBankRow(parsed.row_no, "quarantine", parsed.raw, None, "TERMINAL_SEQ_CONFLICT"))
            self._flag_revision(conn, f"CARD:{existing['id']}", existing["id"], existing["content_hash"], transaction.content_hash)
            return
        now = utc_now()
        conn.execute(
            "INSERT INTO card_transaction(id,terminal_id,seq_id,transaction_type,occurred_at,server_at,posted_date,amount_minor,currency_code,cashback_minor,tip_minor,dcc_minor,arn,authorization_code,masked_card,variable_symbol,variable_symbol_2,issuer,read_method,merchant_place,merchant_address,status,raw_json,content_hash,first_seen_utc,last_seen_utc,row_version) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)",
            (
                uuid4().hex,
                transaction.terminal_id,
                transaction.seq_id,
                transaction.transaction_type,
                transaction.occurred_at.isoformat(sep=" "),
                transaction.server_at.isoformat(sep=" ") if transaction.server_at else None,
                transaction.posted_date.isoformat() if transaction.posted_date else None,
                transaction.amount_minor,
                transaction.currency,
                transaction.cashback_minor,
                transaction.tip_minor,
                transaction.dcc_minor,
                transaction.arn,
                transaction.authorization_code,
                transaction.masked_card,
                transaction.variable_symbol,
                transaction.variable_symbol_2,
                normalize_text(parsed.raw["Vydavatel karty"]) or None,
                normalize_text(parsed.raw["Způsob načtení karty"]) or None,
                normalize_text(parsed.raw["Obchodní místo"]) or None,
                normalize_text(parsed.raw["Adresa obchodního místa"]) or None,
                transaction.status.value,
                canonical_json(parsed.raw),
                transaction.content_hash,
                now,
                now,
            ),
        )
        report.new_rows += 1
        report.sales += 1
        report.totals_minor[transaction.currency] = report.totals_minor.get(transaction.currency, 0) + transaction.amount_minor

    @staticmethod
    def _quarantine(conn: Any, run_id: str, parsed: ParsedBankRow) -> None:
        conn.execute(
            "INSERT INTO quarantined_source_row(id,run_type,run_id,row_no,raw_json,reason_code,reason_text,resolution_json,state,created_at_utc,updated_at_utc) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
            (
                uuid4().hex,
                "BANK",
                run_id,
                parsed.row_no,
                canonical_json(parsed.raw),
                parsed.reason or "UNKNOWN",
                parsed.reason or "UNKNOWN",
                None,
                "OPEN",
                utc_now(),
                utc_now(),
            ),
        )

    @staticmethod
    def _flag_revision(conn: Any, object_ref: str, entity_id: str, old_hash: str, new_hash: str) -> None:
        affected = [row[0] for row in conn.execute(
            "SELECT DISTINCT group_id FROM match_group_source WHERE entity_type='CARD' AND entity_id=? AND active=1", (entity_id,)
        ).fetchall()]
        if not affected:
            return
        now = utc_now()
        conn.execute(
            "INSERT INTO source_revision_alert(id,object_ref,previous_hash,current_hash,changed_fields_json,affected_group_ids_json,state,detected_at_utc) VALUES(?,?,?,?,?,?,?,?)",
            (uuid4().hex, object_ref, old_hash, new_hash, canonical_json(["source_content"]), canonical_json(affected), "OPEN", now),
        )
        placeholders = ",".join("?" for _ in affected)
        conn.execute(
            f"UPDATE match_group SET status='REVIEW_REQUIRED',review_required_reason='Zdrojová data se změnila',row_version=row_version+1,updated_at_utc=? WHERE id IN ({placeholders})",
            (now, *affected),
        )
